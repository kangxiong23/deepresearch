# Layer: Adapter
# File: app/adapters/file_parsers.py
# Responsibility: 实现 FileParserProtocol——从各类文件中提取纯文本内容。
#                 每种文件格式对应一个解析器类，由 FileService 按 can_parse() 路由。
# Input:  文件路径（str）
# Output: 提取的纯文本（str）；无法解析时返回一句**可见的说明**而不是空串
# 禁止: 业务逻辑、导入 UI 库

from __future__ import annotations

import json
import os
import re


# ──────────────────────────────────────────────
# 纯文本 / Markdown
# ──────────────────────────────────────────────

class PlainTextParser:
    """处理 .txt / .md / .rst 等纯文本文件。"""

    _EXTENSIONS = {".txt", ".md", ".markdown", ".rst", ".text"}

    def can_parse(self, filename: str) -> bool:
        return _ext(filename) in self._EXTENSIONS

    def extract_text_sync(self, file_path: str) -> str:
        """同步版本：供上下文构建等同步路径复用，语义与 extract_text 相同。"""
        if not file_path or not os.path.exists(file_path):
            print(f"[FileParser] 文件不存在或路径为空: {repr(file_path)}")
            return f"[无法读取「{_basename(file_path) or file_path}」：文件不存在]"
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        except OSError as e:
            print(f"[FileParser] 读取失败: {file_path} -> {e}")
            return f"[无法读取「{_basename(file_path)}」：{e.strerror or '未知错误'}]"

    async def extract_text(self, file_path: str) -> str:
        return self.extract_text_sync(file_path)


# ──────────────────────────────────────────────
# JSON
# ──────────────────────────────────────────────

class JsonParser:
    """处理 .json 文件，格式化后作为文本注入。"""

    def can_parse(self, filename: str) -> bool:
        return _ext(filename) == ".json"

    def extract_text_sync(self, file_path: str) -> str:
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                data = json.load(f)
            return json.dumps(data, ensure_ascii=False, indent=2)
        except OSError as e:
            return f"[无法读取「{_basename(file_path)}」：{e.strerror or '未知错误'}]"
        except json.JSONDecodeError as e:
            return f"[无法解析「{_basename(file_path)}」：JSON 格式错误（第 {e.lineno} 行 {e.msg}）]"

    async def extract_text(self, file_path: str) -> str:
        return self.extract_text_sync(file_path)


# ──────────────────────────────────────────────
# Word (.docx)
# ──────────────────────────────────────────────

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# 段落遍历时整棵跳过的子树：
#   del         —— 修订模式里被删除的内容，已不是当前稿件的文字
#   txbxContent —— 文本框，另行渲染（混进正文会打乱阅读顺序）
#   instrText   —— 域代码（PAGE / REF 等），要的是域结果里的 w:t
_SKIP_SUBTREE = {"del", "txbxContent", "instrText"}

# 标题判定要同时看 styleId 和样式名：Word 把内置样式 id 存成 "2"/"3" 这种
# 数字（实测真实文档就是这样），而样式名是本地化过的 "Heading 1"/"标题 1"。
# 只认其中一个都会漏。
_HEADING_STYLE_ID = re.compile(r"^Heading(\d)$", re.I)
_HEADING_STYLE_NAME = re.compile(r"^(?:Heading|标题)\s*(\d)$", re.I)

# 页脚常见「第 1 页 共 9 页」「- 3 -」这类纯页码：不是内容，注入只会稀释上下文。
_PAGE_NUMBER_ONLY = re.compile(r"^[\s第页共之\d\-—–‑/·.()（）]*$")

_MAX_TABLE_DEPTH = 2


class DocxParser:
    """
    处理 .docx：按**文档顺序**提取段落、表格、标题、列表、文本框、页眉页脚。

    三条约束（都是踩过的坑）：

    1. **顺序不能丢**。把 doc.paragraphs 与 doc.tables 分两趟跑，等于把文档里
       所有表格都挪到文末 —— 一份"说明 + 表格 + 结论"的文档会变成
       "说明 + 结论 + 表格"。必须走同一趟遍历。

    2. **文字不能少**。Paragraph.text 只认段落的直接子 w:r，因此：
       修订模式里的新增文字（w:ins）、行内内容控件（w:sdt）、块级内容控件
       里的整段，全部读不到；而且 3000 段要 377 ms。改为自己递归遍历
       w:t / w:tab / w:br，实测同等内容 25 ms，且上述三类文字都能取到。

    3. **失败要看得见**。损坏 / 加密 / 改了扩展名的文档返回一句可读说明，
       而不是空串 —— 空串会被上层当成"空文档"静默丢掉，用户只看到
       "传了文件但模型没反应"。

    刻意不做：版式（字体/颜色/缩进）、编号系统的真实还原（numbering.xml）、
    图片里的文字（那是 OCR / 视觉通道的事）。
    """

    def can_parse(self, filename: str) -> bool:
        return _ext(filename) == ".docx"

    def extract_text_sync(self, file_path: str) -> str:
        try:
            import docx  # python-docx
        except ImportError:
            return "[无法解析 .docx：未安装 python-docx]"

        name = _basename(file_path)
        try:
            document = docx.Document(file_path)
        except Exception as e:
            print(f"[FileParser] 打开 docx 失败: {file_path} -> {e}")
            return f"[无法解析「{name}」：{_docx_open_reason(e)}]"

        try:
            return self._render(document, file_path)
        except Exception as e:
            print(f"[FileParser] 解析 docx 失败: {file_path} -> {e}")
            return f"[无法解析「{name}」：文档结构异常（{type(e).__name__}）]"

    async def extract_text(self, file_path: str) -> str:
        return self.extract_text_sync(file_path)

    def extract_images_sync(self, file_path: str) -> list[bytes]:
        """
        提取文档内嵌图片的原始字节（word/media/*），按出现顺序。

        刻意**不做**格式判断与缩放：哪些能发、发多大是 Core 的事
        （app/core/image_attachments.py），适配器只负责把字节拿出来。
        """
        try:
            import zipfile

            with zipfile.ZipFile(file_path) as archive:
                names = [n for n in archive.namelist() if n.startswith("word/media/")]
                names.sort(key=_natural_key)
                return [archive.read(n) for n in names if not n.endswith("/")]
        except Exception as e:
            print(f"[FileParser] 提取 docx 图片失败: {file_path} -> {e}")
            return []

    # ──────────────────────────────────────────
    # 渲染
    # ──────────────────────────────────────────

    def _render(self, document, file_path: str) -> str:
        chunks: list[str] = []
        state = {
            "list": {},
            "images": 0,
            # numId → {ilvl: numFmt}：决定项目符号还是编号（见 _load_numbering）
            "numbering": _load_numbering(file_path),
            # 脚注/尾注正文单独存在 footnotes.xml / endnotes.xml 里，
            # 正文只会留下一个 w:footnoteReference 空标记，必须二次拼回来
            "footnotes": _load_note_entries(file_path, "word/footnotes.xml", "footnote"),
            "endnotes": _load_note_entries(file_path, "word/endnotes.xml", "endnote"),
        }

        for kind, block in self._iter_blocks(document.element.body, document):
            if kind == "p":
                chunks.extend(self._paragraph_chunks(block, state))
            else:
                chunks.extend(self._table_chunks(block, 0))

        chunks.extend(self._header_footer_notes(document))
        chunks.extend(_note_chunks(state["footnotes"], "脚注"))
        chunks.extend(_note_chunks(state["endnotes"], "尾注"))

        media = _count_media(file_path)
        if media:
            chunks.append(
                f"[备注] 文档含 {media} 张图片，未做文字识别"
                "（图片里的文字需另附图片或手动摘录）"
            )

        return "\n\n".join(c for c in chunks if c and c.strip())

    @staticmethod
    def _iter_blocks(container, parent):
        """
        按文档顺序产出 ("p", Paragraph) / ("tbl", Table)。

        展开块级内容控件（w:sdt）—— Word 的封面、目录、以及"文本内容控件"
        都用它包着，Document.iter_inner_content() 会整块跳过（实测）。
        """
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        for child in container.iterchildren():
            tag = child.tag.split("}")[-1]
            if tag == "p":
                yield "p", Paragraph(child, parent)
            elif tag == "tbl":
                yield "tbl", Table(child, parent)
            elif tag == "sdt":
                content = child.find(_W + "sdtContent")
                if content is not None:
                    yield from DocxParser._iter_blocks(content, parent)

    def _paragraph_chunks(self, paragraph, state: dict) -> list[str]:
        """一个段落的输出：正文（带标题/列表前缀）、图片占位、文本框。"""
        element = paragraph._p
        chunks: list[str] = []

        text = _para_text(element)
        if text.strip():
            chunks.append(_format_paragraph(paragraph, text, state))
        else:
            state["list"].clear()

        drawings = _count_drawings(element)
        for _ in range(drawings):
            state["images"] += 1
            chunks.append(f"[图片 {state['images']}]")

        for box in element.findall(".//" + _W + "txbxContent"):
            box_texts = [_squash(_para_text(p)) for p in box.findall(_W + "p")]
            joined = " ".join(t for t in box_texts if t)
            if joined:
                chunks.append("[文本框] " + joined)

        return chunks

    def _table_chunks(self, table, depth: int) -> list[str]:
        """表格渲染成 Markdown 管道表；单元格里的嵌套表格另起一块。"""
        rows, nested = self._table_rows(table, depth)
        if not rows:
            return nested

        width = max(len(r) for r in rows)
        lines: list[str] = []
        for index, row in enumerate(rows):
            padded = row + [""] * (width - len(row))
            lines.append("| " + " | ".join(padded) + " |")
            # 只有 Word 的「重复标题行」标记能可靠说明"这是表头"，别猜
            if index == 0 and _row_is_header(table, 0):
                lines.append("| " + " | ".join(["---"] * width) + " |")

        return ["\n".join(lines)] + nested

    def _table_rows(self, table, depth: int) -> tuple[list[list[str]], list[str]]:
        rows: list[list[str]] = []
        nested: list[str] = []
        for row in table.rows:
            cells: list[str] = []
            last_tc = None
            for cell in row.cells:
                # 跨列合并的单元格会按网格宽度被重复返回同一个 tc
                if cell._tc is last_tc:
                    continue
                last_tc = cell._tc
                text, sub = self._cell_text(cell, depth)
                cells.append(text)
                nested.extend(sub)
            if any(c for c in cells):
                rows.append(cells)
        return rows, nested

    def _cell_text(self, cell, depth: int) -> tuple[str, list[str]]:
        parts: list[str] = []
        for paragraph in cell.paragraphs:
            text = _squash(_para_text(paragraph._p))
            if text:
                parts.append(text)

        nested: list[str] = []
        if depth < _MAX_TABLE_DEPTH:
            for sub in cell.tables:
                chunks = self._table_chunks(sub, depth + 1)
                nested.extend("[嵌套表格]\n" + c for c in chunks)

        return " ".join(parts), nested

    # ──────────────────────────────────────────
    # 页眉 / 页脚
    # ──────────────────────────────────────────

    def _header_footer_notes(self, document) -> list[str]:
        """
        页眉/页脚文本。三种变体都收：普通、首页、偶数页（Word 的"首页不同"
        /"奇偶页不同"），按文本去重——变体没定义时会回落到普通页眉。
        """
        notes: list[str] = []
        seen: set[str] = set()
        try:
            sections = list(document.sections)
        except Exception:
            return notes

        variants = (
            ("页眉", "header"),
            ("页脚", "footer"),
            ("首页页眉", "first_page_header"),
            ("首页页脚", "first_page_footer"),
            ("偶数页页眉", "even_page_header"),
            ("偶数页页脚", "even_page_footer"),
        )

        for section in sections:
            for label, attr in variants:
                try:
                    part = getattr(section, attr)
                    if part is None or part.is_linked_to_previous:
                        continue
                    texts = [_squash(_para_text(p._p)) for p in part.paragraphs]
                    tables = []
                    for table in part.tables:
                        tables.extend(self._table_chunks(table, 0))
                except Exception:
                    continue

                joined = " ".join(t for t in texts if t)
                if tables:
                    joined = " ".join(x for x in [joined, *tables] if x)
                if not joined or joined in seen:
                    continue
                # 纯页码页脚没有信息量，丢掉
                if "页脚" in label and any(ch.isdigit() for ch in joined) \
                        and _PAGE_NUMBER_ONLY.match(joined):
                    continue
                seen.add(joined)
                notes.append(f"[{label}] {joined}")

        return notes


# ──────────────────────────────────────────────
# PDF
# ──────────────────────────────────────────────

class PdfParser:
    """处理 .pdf 文件，逐页提取文本。"""

    def can_parse(self, filename: str) -> bool:
        return _ext(filename) == ".pdf"

    def extract_text_sync(self, file_path: str) -> str:
        try:
            import pdfplumber
        except ImportError:
            return "[无法解析 .pdf：请安装 pdfplumber]"

        name = _basename(file_path)
        try:
            pages: list[str] = []
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        pages.append(text.strip())
        except Exception as e:
            print(f"[FileParser] 解析 pdf 失败: {file_path} -> {e}")
            return f"[无法解析「{name}」：PDF 读取失败（{type(e).__name__}）]"

        if not pages:
            # 扫描件没有文本层：OCR 不在本项目的范围内，但必须说出来
            return f"[「{name}」没有文本层（可能是扫描件，需要 OCR 才能读取）]"
        return "\n\n".join(pages)

    async def extract_text(self, file_path: str) -> str:
        return self.extract_text_sync(file_path)


# ──────────────────────────────────────────────
# 工厂：返回默认解析器列表（注册顺序即优先级）
# ──────────────────────────────────────────────

def default_parsers() -> list:
    """
    返回默认文件解析器列表，供依赖注入时使用。

    使用方式（在 main.py）：
        from app.adapters.file_parsers import default_parsers
        # file_service = FileService(parsers=default_parsers())  # FileService 在 app.core 中
    """
    return [
        PlainTextParser(),
        JsonParser(),
        DocxParser(),
        PdfParser(),
    ]


def unsupported_hints() -> dict[str, str]:
    """
    已知不支持、但有明确处置建议的格式 → 一句人话。

    放在适配器层：文件格式知识属于这里。FileService 只负责把这句话
    拼进"无法解析"的提示里，不自己判断格式。
    """
    return {
        ".doc": "旧版 .doc 格式不支持，请用 Word / WPS 另存为 .docx 后重试",
        ".xls": "暂不支持 .xls，请另存为 .xlsx 或 .csv 后重试",
        ".xlsx": "暂不支持 .xlsx，请另存为 .csv 后重试",
        ".ppt": "暂不支持 .ppt，请另存为 .pdf 后重试",
        ".pptx": "暂不支持 .pptx，请另存为 .pdf 后重试",
        # 图片走"消息附件"那条路（app/core/image_attachments.py），挂载进 system
        # prompt 的图片发不出去——服务端只接受 user 消息里带图。
        ".png": "挂载的图片不会送进模型，请在输入框里作为消息附件发送",
        ".jpg": "挂载的图片不会送进模型，请在输入框里作为消息附件发送",
        ".jpeg": "挂载的图片不会送进模型，请在输入框里作为消息附件发送",
        ".webp": "挂载的图片不会送进模型，请在输入框里作为消息附件发送",
        ".gif": "挂载的图片不会送进模型，请在输入框里作为消息附件发送",
    }


# ──────────────────────────────────────────────
# docx 内部工具
# ──────────────────────────────────────────────

def _para_text(element) -> str:
    """
    段落全文：递归收集 w:t / w:tab / w:br。

    递归（而不是取直接子 w:r）是关键：超链接 w:hyperlink、修订插入 w:ins、
    行内内容控件 w:sdt 都夹在中间一层，Paragraph.text 会漏掉后两者。
    """
    out: list[str] = []
    _walk_text(element, out)
    return "".join(out)


def _walk_text(element, out: list[str]) -> None:
    for child in element:
        tag = child.tag.split("}")[-1]
        if tag in _SKIP_SUBTREE:
            continue
        if tag == "t":
            out.append(child.text or "")
        elif tag == "tab":
            out.append("\t")
        elif tag in ("br", "cr"):
            out.append("\n")
        elif tag == "footnoteReference":
            # 正文里只有一个空引用，脚注正文在 footnotes.xml 里，末尾统一附上
            out.append(f"[脚注 {child.get(_W + 'id') or '?'}]")
        elif tag == "endnoteReference":
            out.append(f"[尾注 {child.get(_W + 'id') or '?'}]")
        else:
            _walk_text(child, out)


def _format_paragraph(paragraph, text: str, state: dict) -> str:
    """按标题层级 / 列表类型给段落加前缀。"""
    heading = _heading_level(paragraph)
    if heading:
        state["list"].clear()
        return "#" * heading + " " + _squash(text)

    kind, level = _list_style(paragraph, state.get("numbering") or {})
    if kind is None:
        state["list"].clear()
        return text.strip()

    indent = "  " * level
    if kind == "bullet":
        return indent + "- " + _squash(text)

    counters = state["list"]
    counters[level] = counters.get(level, 0) + 1
    for deeper in [k for k in counters if k > level]:
        counters.pop(deeper, None)
    return indent + str(counters[level]) + ". " + _squash(text)


def _heading_level(paragraph) -> int:
    """标题层级（1-6），非标题返回 0。"""
    for candidate in (_style_id(paragraph), _style_name(paragraph)):
        if not candidate:
            continue
        matched = _HEADING_STYLE_ID.match(candidate) or _HEADING_STYLE_NAME.match(candidate)
        if matched:
            return min(6, int(matched.group(1)))

    # 段落自己没写大纲级别时，看样式定义里写没写（Word 的标题样式都写了）
    level = _outline_level(paragraph._p.pPr)
    if level:
        return level
    try:
        return _outline_level(paragraph.style.element.pPr)
    except Exception:
        return 0


def _outline_level(pPr) -> int:
    if pPr is None:
        return 0
    outline = pPr.find(_W + "outlineLvl")
    if outline is None:
        return 0
    raw = outline.get(_W + "val")
    if not raw or not raw.isdigit():
        return 0
    depth = int(raw) + 1
    return depth if 1 <= depth <= 6 else 0


def _list_style(paragraph, numbering: dict) -> tuple[str | None, int]:
    """
    (类型, 层级)。类型为 "bullet" / "number"，非列表返回 (None, 0)。

    类型判定顺序：numbering.xml 的 numFmt（最准） → 样式名 → 默认按编号。
    """
    style_id = _style_id(paragraph)
    style_name = _style_name(paragraph)

    pPr = paragraph._p.pPr
    numPr = pPr.numPr if pPr is not None else None

    # 中文版 Word 里「列表段落」样式（ListParagraph）也用于普通缩进，
    # 单靠它判断会把大量正文误判成列表，所以只认确实带编号的样式。
    looks_like_list = (
        style_id.startswith("ListNumber") or style_id.startswith("ListBullet")
        or style_name.startswith(("List Number", "List Bullet", "列表编号", "列表项目"))
    )
    if numPr is None and not looks_like_list:
        return None, 0

    level = 0
    if numPr is not None and numPr.ilvl is not None and numPr.ilvl.val is not None:
        try:
            level = int(numPr.ilvl.val)
        except (TypeError, ValueError):
            level = 0
    level = max(0, min(level, 8))

    kind = _number_format_kind(numPr, numbering, level)
    if kind is None:
        kind = "bullet" if ("Bullet" in style_id or "Bullet" in style_name) else "number"
    return kind, level


def _number_format_kind(numPr, numbering: dict, level: int) -> str | None:
    """按 numbering.xml 的 numFmt 判断类型；查不到返回 None。"""
    if numPr is None or numPr.numId is None or not numbering:
        return None
    try:
        num_id = str(numPr.numId.val)
    except Exception:
        return None
    num_fmt = (numbering.get(num_id) or {}).get(str(level))
    if not num_fmt:
        return None
    # 无编号的列表项（numFmt=none）按项目符号处理，免得被计数成一个假序号
    return "bullet" if num_fmt in ("bullet", "none") else "number"


def _style_id(paragraph) -> str:
    try:
        return paragraph.style.style_id or ""
    except Exception:
        return ""


def _style_name(paragraph) -> str:
    try:
        return paragraph.style.name or ""
    except Exception:
        return ""


def _load_numbering(file_path: str) -> dict:
    """
    读 word/numbering.xml，得到 {numId: {ilvl: numFmt}}。

    只取"格式"（bullet / decimal / chineseCounting …），不还原真实序号值：
    序号值要处理重启编号、多级联动、%1.%2 模板，收益远小于成本。
    失败一律返回空字典（文档没有编号部分、老格式、解析异常都走这条路）。
    """
    root = _read_xml_part(file_path, "word/numbering.xml")
    if root is None:
        return {}

    abstract: dict[str, dict[str, str]] = {}
    for node in root.iter(_W + "abstractNum"):
        abstract_id = node.get(_W + "abstractNumId") or ""
        levels: dict[str, str] = {}
        for lvl in node.findall(_W + "lvl"):
            ilvl = lvl.get(_W + "ilvl") or "0"
            fmt = lvl.find(_W + "numFmt")
            levels[ilvl] = (fmt.get(_W + "val") or "") if fmt is not None else ""
        abstract[abstract_id] = levels

    mapping: dict[str, dict[str, str]] = {}
    for num in root.iter(_W + "num"):
        num_id = num.get(_W + "numId")
        ref = num.find(_W + "abstractNumId")
        if num_id and ref is not None:
            mapping[num_id] = abstract.get(ref.get(_W + "val") or "", {})
    return mapping


def _read_xml_part(file_path: str, part_name: str):
    """
    从 docx 包里读一个原样的 XML 部件（numbering / footnotes / endnotes）。

    python-docx 不暴露这些部件，只能自己解包；任何失败都返回 None，
    调用方按"没有这个部件"处理，不让次要信息把整篇解析拖垮。
    """
    try:
        import zipfile

        from lxml import etree

        with zipfile.ZipFile(file_path) as archive:
            if part_name not in archive.namelist():
                return None
            return etree.fromstring(archive.read(part_name))
    except Exception as e:
        print(f"[FileParser] 读取 {part_name} 失败: {e}")
        return None


def _load_note_entries(file_path: str, part_name: str, tag: str) -> dict[str, str]:
    """
    读脚注/尾注正文：{id: 文本}。

    Word 会在同一个部件里塞几条 w:type="separator" / "continuationSeparator"
    的分隔线条目（id 为 -1 和 0），它们没有正文，必须跳过。
    """
    root = _read_xml_part(file_path, part_name)
    if root is None:
        return {}

    entries: dict[str, str] = {}
    for node in root.findall(_W + tag):
        note_type = node.get(_W + "type") or ""
        if note_type and note_type != "normal":
            continue
        note_id = node.get(_W + "id")
        if note_id is None:
            continue
        texts = [_squash(_para_text(p)) for p in node.findall(_W + "p")]
        joined = " ".join(t for t in texts if t)
        if joined:
            entries[note_id] = joined
    return entries


def _note_chunks(entries: dict[str, str], label: str) -> list[str]:
    """把脚注/尾注按 id 排好附在正文后面，与正文里的 [脚注 N] 标记对得上。"""
    if not entries:
        return []

    def key(item: tuple[str, str]) -> tuple[int, str]:
        raw = item[0]
        try:
            return (int(raw), raw)
        except (TypeError, ValueError):
            return (10 ** 9, raw)

    return [f"[{label} {nid}] {text}" for nid, text in sorted(entries.items(), key=key)]


def _count_drawings(element) -> int:
    """段落里的图片数量（文本框用的 w:pict 不算图片）。"""
    count = 0
    for drawing in element.findall(".//" + _W + "drawing"):
        if drawing.find(".//" + _W + "txbxContent") is None:
            count += 1
    for pict in element.findall(".//" + _W + "pict"):
        if pict.find(".//" + _W + "txbxContent") is None:
            count += 1
    return count


def _row_is_header(table, index: int) -> bool:
    try:
        trPr = table.rows[index]._tr.trPr
    except Exception:
        return False
    if trPr is None:
        return False
    return trPr.find(_W + "tblHeader") is not None


def _count_media(file_path: str) -> int:
    try:
        import zipfile
        with zipfile.ZipFile(file_path) as archive:
            return sum(1 for n in archive.namelist() if n.startswith("word/media/"))
    except Exception:
        return 0


def _docx_open_reason(error: Exception) -> str:
    text = str(error).lower()
    if "package not found" in text or "not a zip" in text or "badzip" in text:
        return "文件已损坏，或只是把 .doc 改了扩展名（请用 Word 另存为 .docx）"
    if "encrypt" in text or "password" in text:
        return "文档被加密，需要先解除密码保护"
    return f"无法打开（{type(error).__name__}）"


def _squash(text: str) -> str:
    """把换行/制表/连续空格压成单空格（标题、列表、表格单元格用）。"""
    return " ".join((text or "").split())


# ──────────────────────────────────────────────
# 工具
# ──────────────────────────────────────────────

def _ext(filename: str) -> str:
    """提取小写扩展名，如 'report.PDF' → '.pdf'。"""
    return os.path.splitext(filename or "")[1].lower()


def _basename(path: str) -> str:
    return os.path.basename(path or "") or (path or "")


def _natural_key(name: str) -> tuple:
    """让 image2.png 排在 image10.png 前面（Word 按插入顺序编号）。"""
    parts = re.split(r"(\d+)", name)
    return tuple((0, int(p)) if p.isdigit() else (1, p) for p in parts)
