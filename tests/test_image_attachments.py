# -*- coding: utf-8 -*-
"""
图片输入（视觉）的单元测试：格式嗅探、缩放、token 记账、上下文组装、附件落库。

不访问网络，也不碰真实 data/：blob 仓库用 FileBlobStore(root=临时目录)。
"""

import asyncio
import base64
import io
import os
import shutil
import unittest
import uuid

from PIL import Image

from config import ConfigScope
from app.core import image_attachments, message_attachments
from app.core.context_service import ContextService
from app.core.conversation_service import ConversationService
from app.core.file_service import FileService
from app.adapters.file_parsers import default_parsers
from app.storage.blob_store import FileBlobStore
from app.storage.models import Message, Role
from datetime import datetime

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_FIXTURE_ROOT = os.path.join(_ROOT, "logs", "_test_vision")


def png_bytes(width=64, height=48, color=(220, 20, 20), fmt="PNG") -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buffer, format=fmt)
    return buffer.getvalue()


class _Tmp(unittest.TestCase):
    def setUp(self):
        os.makedirs(_FIXTURE_ROOT, exist_ok=True)
        self.dir = os.path.join(_FIXTURE_ROOT, uuid.uuid4().hex)
        os.makedirs(self.dir)
        self.blobs = FileBlobStore(root=self.dir)

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def write(self, name: str, data: bytes) -> str:
        path = os.path.join(self.dir, name)
        with open(path, "wb") as handle:
            handle.write(data)
        return path


class TestSniffing(unittest.TestCase):
    def test_recognises_supported_formats(self):
        self.assertEqual(image_attachments.sniff_mime(png_bytes(fmt="PNG")), "image/png")
        self.assertEqual(image_attachments.sniff_mime(png_bytes(fmt="JPEG")), "image/jpeg")
        self.assertEqual(image_attachments.sniff_mime(png_bytes(fmt="GIF")), "image/gif")
        self.assertEqual(image_attachments.sniff_mime(png_bytes(fmt="WEBP")), "image/webp")

    def test_rejects_other_content(self):
        self.assertIsNone(image_attachments.sniff_mime(b"%PDF-1.7 ..."))
        self.assertIsNone(image_attachments.sniff_mime(b""))
        self.assertIsNone(image_attachments.sniff_mime(b"PK\x03\x04docx"))

    def test_only_flash_sees_images(self):
        """实测：pro 收到图片会静默丢弃（prompt_tokens 里没有图片）。"""
        self.assertTrue(image_attachments.supports_images("deepseek-v4-flash"))
        self.assertTrue(image_attachments.supports_images("deepseek-flash"))
        self.assertFalse(image_attachments.supports_images("deepseek-v4-pro"))


class TestPrepare(unittest.TestCase):
    def test_small_image_passes_through_unchanged(self):
        data = png_bytes(200, 150)
        prepared = image_attachments.prepare(data)
        self.assertEqual(prepared["data"], data)
        self.assertEqual(prepared["mime"], "image/png")
        self.assertEqual((prepared["width"], prepared["height"]), (200, 150))

    def test_large_image_is_downscaled(self):
        data = png_bytes(3000, 2000)
        prepared = image_attachments.prepare(data)
        self.assertIsNotNone(prepared)
        pixels = prepared["width"] * prepared["height"]
        self.assertLessEqual(pixels, image_attachments.MAX_PIXELS * 1.02)
        self.assertLessEqual(len(prepared["data"]), image_attachments.MAX_IMAGE_BYTES)

    def test_non_image_returns_none(self):
        self.assertIsNone(image_attachments.prepare(b"hello world"))

    def test_token_estimate_bounds(self):
        self.assertGreaterEqual(image_attachments.estimate_tokens(1, 1), 96)
        self.assertLessEqual(image_attachments.estimate_tokens(5000, 5000), 1024)
        self.assertGreater(
            image_attachments.estimate_tokens(1600, 1200),
            image_attachments.estimate_tokens(400, 300),
        )


class TestMessageAttachments(unittest.TestCase):
    def test_image_attachment_renders_marker(self):
        text = message_attachments.expand(
            "看图", [{"name": "shot.png", "kind": "image", "blob_id": "a" * 64}]
        )
        self.assertIn("[图片：shot.png]", text)

    def test_marker_can_explain_why_not_sent(self):
        text = message_attachments.expand(
            "看图",
            [{"name": "shot.png", "kind": "image", "blob_id": "a" * 64}],
            image_suffix="（当前模型不支持图片，未发送）",
        )
        self.assertIn("未发送", text)

    def test_image_tokens_are_counted(self):
        plain = message_attachments.estimate_tokens("看这张图")
        with_image = message_attachments.estimate_tokens(
            "看这张图",
            [{"name": "s.png", "kind": "image", "blob_id": "a" * 64,
              "width": 1600, "height": 1200}],
        )
        self.assertGreaterEqual(with_image - plain, 200)


class _StubStore:
    def get_blocks(self):
        return []

    def list_templates(self):
        return []


def _make_service(blobs, history):
    service = ContextService(
        message_repo=None,
        context_store=_StubStore(),
        tree_store=None,
        blob_store=blobs,
    )
    service.get_enabled_history_messages = lambda **kw: list(history)
    return service


def _store_image(blobs, name="shot.png", size=(64, 48)):
    prepared = image_attachments.prepare(png_bytes(*size))
    record = blobs.put_bytes(prepared["data"], name=name)
    return {
        "name": name, "kind": "image", "blob_id": record["blob_id"],
        "mime": prepared["mime"], "width": prepared["width"], "height": prepared["height"],
    }


class TestContextImages(_Tmp):
    def test_new_message_carries_image_block(self):
        with ConfigScope(model_type="deepseek-v4-flash"):
            image = _store_image(self.blobs)
            context = _make_service(self.blobs, []).build_llm_context(
                "c-1", "这张图是什么", new_user_attachments=[image]
            )

        content = context.messages[-1]["content"]
        self.assertIsInstance(content, list)
        self.assertEqual(content[0]["type"], "text")
        self.assertIn("[图片：shot.png]", content[0]["text"])
        self.assertEqual(content[1]["type"], "image_url")
        self.assertTrue(content[1]["image_url"]["url"].startswith("data:image/png;base64,"))

    def test_image_count_is_capped(self):
        with ConfigScope(model_type="deepseek-v4-flash"):
            images = [_store_image(self.blobs, "s%d.png" % i) for i in range(6)]
            context = _make_service(self.blobs, []).build_llm_context(
                "c-1", "看图", new_user_attachments=images
            )

        blocks = context.messages[-1]["content"]
        self.assertEqual(sum(1 for b in blocks if b["type"] == "image_url"),
                         image_attachments.MAX_IMAGES_PER_REQUEST)

    def test_model_without_vision_degrades_to_text(self):
        with ConfigScope(model_type="deepseek-v4-pro"):
            image = _store_image(self.blobs)
            context = _make_service(self.blobs, []).build_llm_context(
                "c-1", "这张图是什么", new_user_attachments=[image]
            )

        content = context.messages[-1]["content"]
        self.assertIsInstance(content, str)
        self.assertIn("未发送", content)

    def test_sticky_history_message_keeps_its_image(self):
        with ConfigScope(model_type="deepseek-v4-flash"):
            image = _store_image(self.blobs)
            history = [
                Message(id="m-old", conversation_id="c-1", role=Role.USER,
                        content="更早的问题", created_at=datetime.now(),
                        attachments=[{"name": "old.png", "kind": "image",
                                      "blob_id": "b" * 64, "mime": "image/png"}]),
                Message(id="m-new", conversation_id="c-1", role=Role.USER,
                        content="看这张图", created_at=datetime.now(),
                        attachments=[image]),
                Message(id="m-a", conversation_id="c-1", role=Role.ASSISTANT,
                        content="好的", created_at=datetime.now()),
            ]
            context = _make_service(self.blobs, history).build_llm_context(
                "c-1", "那第二点呢"
            )

        # 最近一条带图的历史消息继续带图，更早那条退化成标注
        self.assertIsInstance(context.messages[1]["content"], list)
        self.assertIn("[图片：old.png]", context.messages[0]["content"])
        self.assertIsInstance(context.messages[2]["content"], str)

    def test_missing_blob_is_reported_in_text(self):
        with ConfigScope(model_type="deepseek-v4-flash"):
            ghost = {"name": "ghost.png", "kind": "image", "blob_id": "c" * 64,
                     "mime": "image/png", "width": 10, "height": 10}
            context = _make_service(self.blobs, []).build_llm_context(
                "c-1", "看图", new_user_attachments=[ghost]
            )

        content = context.messages[-1]["content"]
        self.assertIsInstance(content, str)
        self.assertIn("图片已丢失", content)


class TestExtractAttachments(_Tmp):
    def _service(self):
        service = ConversationService.__new__(ConversationService)
        service._file_svc = FileService(parsers=default_parsers())
        service._blobs = self.blobs
        return service

    def test_image_file_becomes_blob_attachment(self):
        path = self.write("shot.png", png_bytes(120, 90))

        attachments = asyncio.run(self._service()._extract_attachments([path]))

        self.assertEqual(len(attachments), 1)
        entry = attachments[0]
        self.assertEqual(entry["kind"], "image")
        self.assertTrue(self.blobs.exists(entry["blob_id"]))
        self.assertEqual((entry["width"], entry["height"]), (120, 90))

    def test_non_image_still_parsed_as_text(self):
        path = self.write("note.txt", "预算 245 万元".encode("utf-8"))

        attachments = asyncio.run(self._service()._extract_attachments([path]))

        self.assertEqual(attachments[0]["text"], "预算 245 万元")
        self.assertNotEqual(attachments[0].get("kind"), "image")

    def test_docx_embedded_image_is_sent(self):
        from docx import Document

        document = Document()
        document.add_paragraph("正文")
        document.add_picture(self.write("pic.png", png_bytes(80, 60)))
        path = os.path.join(self.dir, "with-image.docx")
        document.save(path)

        attachments = asyncio.run(self._service()._extract_attachments([path]))

        kinds = [a.get("kind") for a in attachments]
        self.assertIn("image", kinds)
        text_entry = next(a for a in attachments if a.get("kind") != "image")
        self.assertIn("正文", text_entry["text"])

    def test_too_many_images_are_capped(self):
        paths = [self.write("p%d.png" % i, png_bytes(32, 32)) for i in range(7)]

        attachments = asyncio.run(self._service()._extract_attachments(paths))

        images = [a for a in attachments if a.get("kind") == "image"]
        self.assertEqual(len(images), image_attachments.MAX_IMAGES_PER_REQUEST)


class TestDocxImages(_Tmp):
    def test_extract_images_returns_media_bytes(self):
        from docx import Document
        from app.adapters.file_parsers import DocxParser

        document = Document()
        document.add_picture(self.write("pic.png", png_bytes(80, 60)))
        path = os.path.join(self.dir, "imgs.docx")
        document.save(path)

        images = DocxParser().extract_images_sync(path)

        self.assertEqual(len(images), 1)
        self.assertEqual(image_attachments.sniff_mime(images[0]), "image/png")


class TestPayload(unittest.TestCase):
    def test_block_content_survives_payload_build(self):
        from app.adapters.deepseek_client import DeepSeekClient
        from app.storage.models import LLMContext

        blocks = [
            {"type": "text", "text": "看图"},
            {"type": "image_url", "image_url": {"url": "data:image/png;base64,AAAA"}},
        ]
        context = LLMContext(
            messages=[{"role": "user", "content": blocks}],
            system_prompt="系统提示",
            model_type="deepseek-v4-flash",
            thinking_enabled=False,
            max_tokens=100,
            temperature=0.7,
        )

        payload = DeepSeekClient._build_payload(context)

        self.assertEqual(payload["messages"][1]["content"], blocks)


if __name__ == "__main__":
    unittest.main()
