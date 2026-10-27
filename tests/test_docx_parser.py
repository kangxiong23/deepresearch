# -*- coding: utf-8 -*-
"""
DocxParser 的结构保真与"失败可见"测试。

fixture 全部**程序生成**（不往仓库里塞二进制 docx），生成目录放在 logs/ 下：
tests/base.py 的 tempfile.mkdtemp() 在本机不可写，所以这里显式指定 dir。
"""

import asyncio
import os
import shutil
import unittest
import uuid
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.shared import Inches

from app.adapters.file_parsers import DocxParser, PlainTextParser, unsupported_hints
from app.core.file_service import FileService

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_FIXTURE_ROOT = os.path.join(_ROOT, "logs", "_test_docx")

# 1x1 透明 PNG
_TINY_PNG = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)
import base64


class _DocxFixture(unittest.TestCase):
    def setUp(self):
        # 刻意不用 tempfile.mkdtemp：它以 0700 建目录，在 Windows 上落成
        # 只允许属主访问的 ACL，沙箱内的写入会被拒绝。
        os.makedirs(_FIXTURE_ROOT, exist_ok=True)
        self.dir = os.path.join(_FIXTURE_ROOT, uuid.uuid4().hex)
        os.makedirs(self.dir)
        self.parser = DocxParser()

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    # ── 工具 ──────────────────────────────────
    def _path(self, name="sample.docx"):
        return os.path.join(self.dir, name)

    def _save(self, document, name="sample.docx"):
        path = self._path(name)
        document.save(path)
        return path

    def _png(self):
        path = os.path.join(self.dir, "tiny.png")
        with open(path, "wb") as f:
            f.write(base64.b64decode(_TINY_PNG))
        return path

    def _xml(self, fragment):
        """fragment 由调用方自行插值（有的片段要插两个命名空间）。"""
        return parse_xml(fragment)


class TestDocxStructure(_DocxFixture):
    """① 结构保真：顺序、合并单元格、修订、内容控件、文本框、表格。"""

    def test_table_stays_in_document_order(self):
        document = Document()
        document.add_paragraph("第一段")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "甲"
        table.cell(0, 1).text = "乙"
        document.add_paragraph("第三段")

        text = self.parser.extract_text_sync(self._save(document))

        self.assertLess(text.index("第一段"), text.index("| 甲 | 乙 |"))
        self.assertLess(text.index("| 甲 | 乙 |"), text.index("第三段"))

    def test_merged_cell_not_duplicated(self):
        document = Document()
        table = document.add_table(rows=1, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = "合并"
        table.cell(0, 2).text = "尾巴"

        text = self.parser.extract_text_sync(self._save(document))

        self.assertEqual(text.count("合并"), 1)
        self.assertIn("| 合并 | 尾巴 |", text)

    def test_cell_newline_kept_on_one_row(self):
        document = Document()
        table = document.add_table(rows=1, cols=2)
        cell = table.cell(0, 0)
        cell.text = "第一行"
        cell.add_paragraph("第二行")

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("| 第一行 第二行 |", text)

    def test_tracked_insertion_kept_and_deletion_dropped(self):
        document = Document()
        paragraph = document.add_paragraph("原有文字")
        paragraph._p.append(self._xml(
            '<w:ins xmlns:w="%s" w:id="7" w:author="x" w:date="2020-01-01T00:00:00Z">'
            '<w:r><w:t>新增文字</w:t></w:r></w:ins>' % _W
        ))
        paragraph._p.append(self._xml(
            '<w:del xmlns:w="%s" w:id="8" w:author="x" w:date="2020-01-01T00:00:00Z">'
            '<w:r><w:delText>已删除文字</w:delText></w:r></w:del>' % _W
        ))

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("原有文字新增文字", text)
        self.assertNotIn("已删除文字", text)

    def test_block_content_control_paragraph_extracted(self):
        document = Document()
        document.add_paragraph("正文一")
        document.element.body.insert(1, self._xml(
            '<w:sdt xmlns:w="%s"><w:sdtPr/><w:sdtContent>'
            '<w:p><w:r><w:t>控件里的段落</w:t></w:r></w:p>'
            '</w:sdtContent></w:sdt>' % _W
        ))

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("控件里的段落", text)

    def test_inline_content_control_extracted(self):
        document = Document()
        paragraph = document.add_paragraph("段落开头")
        paragraph._p.append(self._xml(
            '<w:sdt xmlns:w="%s"><w:sdtPr/><w:sdtContent>'
            '<w:r><w:t>行内控件</w:t></w:r></w:sdtContent></w:sdt>' % _W
        ))

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("段落开头行内控件", text)

    def test_hyperlink_text_kept(self):
        document = Document()
        paragraph = document.add_paragraph("参见 ")
        paragraph._p.append(self._xml(
            '<w:hyperlink xmlns:w="%s"'
            ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
            ' r:id="rId9"><w:r><w:t>官方网站</w:t></w:r></w:hyperlink>' % _W
        ))

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("参见 官方网站", text)

    def test_textbox_rendered_once(self):
        document = Document()
        document.add_paragraph("正文")
        element = self._xml(
            '<w:p xmlns:w="%s" xmlns:v="urn:schemas-microsoft-com:vml">'
            '<w:r><w:pict><v:shape id="tb1" style="width:200pt;height:50pt">'
            '<v:textbox><w:txbxContent><w:p><w:r><w:t>框里的字</w:t></w:r></w:p>'
            '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>' % _W
        )
        document.element.body.insert(1, element)

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("[文本框] 框里的字", text)
        self.assertEqual(text.count("框里的字"), 1)

    def test_nested_table_content_extracted(self):
        document = Document()
        table = document.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        cell.text = "外层"
        inner = cell.add_table(rows=1, cols=2)
        inner.cell(0, 0).text = "内左"
        inner.cell(0, 1).text = "内右"

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("外层", text)
        self.assertIn("内左", text)
        self.assertIn("内右", text)

    def test_image_placeholder_and_note(self):
        document = Document()
        document.add_paragraph("下图：")
        document.add_picture(self._png(), width=Inches(0.2))

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("[图片 1]", text)
        self.assertIn("未做文字识别", text)


class TestDocxMarkup(_DocxFixture):
    """标题层级与列表标记。"""

    def test_heading_levels(self):
        document = Document()
        document.add_heading("一级标题", level=1)
        document.add_heading("二级标题", level=2)
        document.add_paragraph("正文")

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("# 一级标题", text)
        self.assertIn("## 二级标题", text)
        self.assertNotIn("# 正文", text)

    def test_numbered_list_gets_counters(self):
        document = Document()
        document.add_paragraph("项一", style="List Number")
        document.add_paragraph("项二", style="List Number")
        document.add_paragraph("普通段落")
        document.add_paragraph("重新计数", style="List Number")

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("1. 项一", text)
        self.assertIn("2. 项二", text)
        self.assertIn("1. 重新计数", text)

    def test_bullet_list_gets_dash(self):
        document = Document()
        document.add_paragraph("要点", style="List Bullet")

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("- 要点", text)

    def _numbering_paragraph(self, document, num_id, text):
        """给段落挂上 numPr，模拟 Word 的真实自动编号（而不是样式自带编号）。"""
        paragraph = document.add_paragraph(text)
        paragraph._p.get_or_add_pPr().append(self._xml(
            '<w:numPr xmlns:w="%s"><w:ilvl w:val="0"/><w:numId w:val="%s"/></w:numPr>'
            % (_W, num_id)
        ))
        return paragraph

    def test_bullet_numbering_format_renders_dash(self):
        document = Document()
        self._numbering_paragraph(document, 1, "符号项")   # 默认模板 numId=1 是 bullet

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("- 符号项", text)

    def test_decimal_numbering_format_gets_counter(self):
        document = Document()
        self._numbering_paragraph(document, 5, "第一项")   # 默认模板 numId=5 是 decimal
        self._numbering_paragraph(document, 5, "第二项")

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("1. 第一项", text)
        self.assertIn("2. 第二项", text)


class TestDocxHeaderFooter(_DocxFixture):
    def test_header_and_footer_captured(self):
        document = Document()
        document.add_paragraph("正文")
        section = document.sections[0]
        section.header.paragraphs[0].text = "内部资料请勿外传"
        section.footer.paragraphs[0].text = "版本 v3"

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("[页眉] 内部资料请勿外传", text)
        self.assertIn("[页脚] 版本 v3", text)

    def test_page_number_footer_dropped(self):
        document = Document()
        document.add_paragraph("正文")
        section = document.sections[0]
        section.header.paragraphs[0].text = "项目周报"
        section.footer.paragraphs[0].text = "第 1 页 共 9 页"

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("[页眉] 项目周报", text)
        self.assertNotIn("第 1 页", text)


class TestDocxNotesAndVariants(_DocxFixture):
    """② 覆盖补全：脚注/尾注、页眉页脚的多节变体与表格。"""

    def _inject_part(self, path, part_name, xml):
        """把额外的 XML 部件塞进已保存的 docx（python-docx 没有脚注 API）。"""
        import zipfile

        with zipfile.ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        entries[part_name] = xml.encode("utf-8")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in entries.items():
                archive.writestr(name, data)

    def _notes_xml(self, wrapper, tag, entries):
        body = "".join(
            '<w:%s w:id="%s"><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:%s>'
            % (tag, note_id, text, tag)
            for note_id, text in entries
        )
        return (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:%s xmlns:w="%s">'
            '<w:%s w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:%s>'
            '<w:%s w:type="continuationSeparator" w:id="0">'
            '<w:p><w:r><w:continuationSeparator/></w:r></w:p></w:%s>'
            '%s</w:%s>' % (wrapper, _W, tag, tag, tag, tag, body, wrapper)
        )

    def _reference_run(self, reference_tag, note_id):
        return self._xml(
            '<w:r xmlns:w="%s"><w:%s w:id="%s"/></w:r>' % (_W, reference_tag, note_id)
        )

    def test_footnotes_marked_inline_and_listed(self):
        document = Document()
        paragraph = document.add_paragraph("正文带脚注")
        paragraph._p.append(self._reference_run("footnoteReference", "1"))
        path = self._save(document)
        self._inject_part(path, "word/footnotes.xml", self._notes_xml(
            "footnotes", "footnote", [("1", "第一条脚注"), ("2", "第二条脚注")]))

        text = self.parser.extract_text_sync(path)

        self.assertIn("正文带脚注[脚注 1]", text)
        self.assertIn("[脚注 1] 第一条脚注", text)
        self.assertIn("[脚注 2] 第二条脚注", text)
        self.assertNotIn("separator", text)

    def test_endnotes_listed(self):
        document = Document()
        paragraph = document.add_paragraph("正文带尾注")
        paragraph._p.append(self._reference_run("endnoteReference", "3"))
        path = self._save(document)
        self._inject_part(path, "word/endnotes.xml", self._notes_xml(
            "endnotes", "endnote", [("3", "一条尾注")]))

        text = self.parser.extract_text_sync(path)

        self.assertIn("[尾注 3]", text)
        self.assertIn("[尾注 3] 一条尾注", text)

    def test_missing_notes_part_is_harmless(self):
        document = Document()
        document.add_paragraph("没有脚注的普通文档")

        text = self.parser.extract_text_sync(self._save(document))

        self.assertNotIn("[脚注", text)
        self.assertIn("没有脚注的普通文档", text)

    def test_header_table_captured(self):
        document = Document()
        document.add_paragraph("正文")
        header = document.sections[0].header
        table = header.add_table(rows=1, cols=2, width=Inches(6))
        table.cell(0, 0).text = "项目"
        table.cell(0, 1).text = "编号"

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("[页眉]", text)
        self.assertIn("| 项目 | 编号 |", text)

    def test_first_page_header_captured(self):
        document = Document()
        document.add_paragraph("正文")
        section = document.sections[0]
        section.header.paragraphs[0].text = "普通页眉"
        section.different_first_page_header_footer = True
        section.first_page_header.paragraphs[0].text = "首页专用页眉"

        text = self.parser.extract_text_sync(self._save(document))

        self.assertIn("[页眉] 普通页眉", text)
        self.assertIn("[首页页眉] 首页专用页眉", text)


class TestDocxFailureVisibility(_DocxFixture):
    """③ 失败可见：坏文件、空文档、不支持的格式都不再返回空串。"""

    def test_corrupt_docx_reports_visible_message(self):
        path = self._path("broken.docx")
        with open(path, "wb") as f:
            f.write(b"PK\x03\x04 not really a docx")

        text = self.parser.extract_text_sync(path)

        self.assertTrue(text.startswith("[无法解析"))
        self.assertIn("损坏", text)

    def test_empty_docx_returns_empty_string(self):
        path = self._save(Document(), "empty.docx")

        self.assertEqual(self.parser.extract_text_sync(path), "")

    def test_empty_docx_becomes_visible_attachment_text(self):
        path = self._save(Document(), "empty.docx")
        service = FileService(parsers=[self.parser])

        results = asyncio.run(service.extract_files([path]))

        self.assertEqual(len(results), 1)
        self.assertIn("没有可提取的文字内容", results[0][1])

    def test_unsupported_extension_uses_hint(self):
        path = self._path("old.doc")
        with open(path, "wb") as f:
            f.write(b"\xd0\xcf\x11\xe0")
        service = FileService(parsers=[self.parser], unsupported_hints=unsupported_hints())

        results = asyncio.run(service.extract_files([path]))

        self.assertIn("旧版 .doc", results[0][1])
        self.assertIn(".docx", results[0][1])

    def test_unknown_extension_still_reported(self):
        path = self._path("archive.zip")
        with open(path, "wb") as f:
            f.write(b"PK")
        service = FileService(parsers=[self.parser])

        results = asyncio.run(service.extract_files([path]))

        self.assertIn("不支持的文件类型", results[0][1])

    def test_parser_exception_is_reported(self):
        class _Boom:
            def can_parse(self, filename):
                return True

            async def extract_text(self, path):
                raise RuntimeError("boom")

        service = FileService(parsers=[_Boom()])

        results = asyncio.run(service.extract_files(["whatever.txt"]))

        self.assertIn("解析过程出错", results[0][1])

    def test_missing_file_reported_as_missing(self):
        service = FileService(parsers=[PlainTextParser()])

        results = asyncio.run(service.extract_files([os.path.join(self.dir, "nope.txt")]))

        self.assertIn("文件不存在", results[0][1])

    def test_no_file_is_ever_dropped(self):
        path = self._path("old.doc")
        with open(path, "wb") as f:
            f.write(b"x")
        service = FileService(parsers=[self.parser])

        results = asyncio.run(service.extract_files([path, "missing-anything.xyz"]))

        self.assertEqual(len(results), 2)


class _FakeBlobs:
    """只实现 read_mount_text 用到的三个方法。"""

    def __init__(self, path):
        self.path = Path(path)
        self.written = []

    def read_cached_text(self, blob_id):
        return None

    def blob_path(self, blob_id):
        return self.path

    def write_cached_text(self, blob_id, text):
        self.written.append(text)


class TestMountCache(_DocxFixture):
    def test_failure_is_not_cached(self):
        path = self._path("broken_blob")
        with open(path, "wb") as f:
            f.write(b"PK\x03\x04 broken")
        blobs = _FakeBlobs(path)
        service = FileService(parsers=[self.parser])

        text = service.read_mount_text(blobs, {"blob_id": "deadbeef", "name": "x.docx"})

        self.assertTrue(text.startswith("[无法解析"))
        self.assertEqual(blobs.written, [])

    def test_success_is_cached(self):
        document = Document()
        document.add_paragraph("能读的内容")
        path = self._save(document, "blob_ok")
        blobs = _FakeBlobs(path)
        service = FileService(parsers=[self.parser])

        text = service.read_mount_text(blobs, {"blob_id": "cafe", "name": "x.docx"})

        self.assertIn("能读的内容", text)
        self.assertEqual(len(blobs.written), 1)


if __name__ == "__main__":
    unittest.main()
