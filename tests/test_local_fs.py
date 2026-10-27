# Layer: tests
# File: tests/test_local_fs.py
# Responsibility: 验证本地文件只读适配器——扫描过滤、编码探测、分页读取、
#                 路径越界校验、docx 提取。

from __future__ import annotations

import json
import os
import shutil
import unittest

from app.adapters import local_fs


class LocalFsTestBase(unittest.TestCase):
    def setUp(self) -> None:
        # 不用 tempfile.mkdtemp：它创建的目录在沙箱下不可写（WinError 5），
        # 只有 makedirs 建的目录可正常写入。
        base = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            "logs", "_test_local_fs",
        )
        self.root = os.path.join(base, self._testMethodName)
        shutil.rmtree(self.root, ignore_errors=True)
        os.makedirs(self.root, exist_ok=True)
        self.addCleanup(shutil.rmtree, self.root, True)

    def write(self, rel: str, data: bytes) -> str:
        path = os.path.join(self.root, rel)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "wb") as f:
            f.write(data)
        return path


class TestScan(LocalFsTestBase):
    def test_only_supported_exts_and_deny_dirs(self) -> None:
        self.write("a.txt", "你好".encode("utf-8"))
        self.write("b.md", "# 标题".encode("utf-8"))
        self.write("c.json", b'{"k": 1}')
        self.write("d.jsonl", b'{"a":1}\n{"a":2}\n')
        self.write("skip.pdf", b"%PDF-1.4")
        self.write("skip.exe", b"MZ")
        self.write(os.path.join("node_modules", "x.md"), b"denied")
        self.write(os.path.join(".git", "config"), b"denied")

        entries, stats = local_fs.scan([self.root])
        names = sorted(e.name for e in entries)
        self.assertEqual(names, ["a.txt", "b.md", "c.json", "d.jsonl"])
        self.assertGreaterEqual(stats.skipped, 2)

    def test_max_files_sets_truncated(self) -> None:
        for i in range(10):
            self.write(f"f{i}.txt", b"x")
        entries, stats = local_fs.scan([self.root], max_files=4)
        self.assertEqual(len(entries), 4)
        self.assertTrue(stats.truncated)

    def test_depth_limit(self) -> None:
        self.write(os.path.join("a", "b", "c", "deep.txt"), b"x")
        entries, _ = local_fs.scan([self.root], max_depth=1)
        self.assertFalse(any(e.name == "deep.txt" for e in entries))

    def test_missing_root_is_counted_not_raised(self) -> None:
        entries, stats = local_fs.scan([os.path.join(self.root, "nope")])
        self.assertEqual(entries, [])
        self.assertGreaterEqual(stats.errors, 1)


class TestEncoding(LocalFsTestBase):
    def test_utf8_bom(self) -> None:
        p = self.write("bom.txt", b"\xef\xbb\xbf" + "中文".encode("utf-8"))
        res = local_fs.read(p)
        self.assertEqual(res.encoding, "utf-8-sig")
        self.assertEqual(res.text, "中文")

    def test_plain_utf8(self) -> None:
        p = self.write("u8.txt", "中文内容\n第二行".encode("utf-8"))
        res = local_fs.read(p)
        self.assertEqual(res.encoding, "utf-8")
        self.assertIn("第二行", res.text)

    def test_gb18030(self) -> None:
        p = self.write("gbk.txt", "这是 GB18030 编码的中文\n第二行".encode("gb18030"))
        res = local_fs.read(p)
        self.assertEqual(res.encoding, "gb18030")
        self.assertIn("这是 GB18030", res.text)
        self.assertTrue(any("gb18030" in w for w in res.warnings))

    def test_utf16le_without_bom(self) -> None:
        p = self.write("u16.txt", "中文\n第二行".encode("utf-16-le"))
        res = local_fs.read(p)
        self.assertTrue(res.encoding.startswith("utf-16"))
        self.assertIn("第二行", res.text)


class TestRead(LocalFsTestBase):
    def test_paging(self) -> None:
        body = "\n".join(f"line{i}" for i in range(1, 501)).encode("utf-8")
        p = self.write("long.txt", body)
        res = local_fs.read(p, start_line=10, max_lines=5)
        self.assertEqual(res.total_lines, 500)
        self.assertEqual(res.start_line, 10)
        self.assertEqual(res.end_line, 14)
        self.assertEqual(res.text.splitlines()[0], "line10")
        self.assertTrue(res.truncated)

    def test_missing_file_returns_error(self) -> None:
        res = local_fs.read(os.path.join(self.root, "nope.txt"))
        self.assertTrue(res.error)

    def test_unsupported_ext_rejected(self) -> None:
        p = self.write("x.pdf", b"%PDF")
        res = local_fs.read(p)
        self.assertIn("不支持", res.error)

    def test_json_pretty_printed(self) -> None:
        p = self.write("d.json", json.dumps({"a": 1, "b": [1, 2]}).encode("utf-8"))
        res = local_fs.read(p)
        self.assertEqual(res.error, "")
        self.assertIn('"a": 1', res.text)

    def test_broken_json_reports_line(self) -> None:
        p = self.write("bad.json", b'{"a": }')
        res = local_fs.read(p)
        self.assertIn("JSON 解析失败", res.error)

    def test_jsonl_reads_by_line(self) -> None:
        p = self.write("d.jsonl", b'{"a":1}\n{"a":2}\n{"a":3}\n')
        res = local_fs.read(p, start_line=2, max_lines=1)
        self.assertEqual(res.text, '{"a":2}')
        self.assertEqual(res.total_lines, 3)

    def test_docx_paragraphs_and_tables(self) -> None:
        import docx

        p = os.path.join(self.root, "doc.docx")
        d = docx.Document()
        d.add_paragraph("第一段")
        d.add_paragraph("第二段")
        table = d.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "A"
        table.rows[0].cells[1].text = "B"
        d.save(p)

        res = local_fs.read(p)
        self.assertEqual(res.error, "")
        self.assertIn("第一段", res.text)
        self.assertIn("A | B", res.text)
        self.assertTrue(any("页眉" in w for w in res.warnings))


class TestContainment(LocalFsTestBase):
    def test_inside(self) -> None:
        p = self.write("sub/a.txt", b"x")
        self.assertTrue(local_fs.is_within(p, [self.root]))

    def test_root_itself(self) -> None:
        self.assertTrue(local_fs.is_within(self.root, [self.root]))

    def test_sibling_prefix_not_confused(self) -> None:
        p = self.write("a.txt", b"x")
        sibling = self.root + "_other"
        os.makedirs(sibling, exist_ok=True)
        self.addCleanup(shutil.rmtree, sibling, True)
        self.assertFalse(local_fs.is_within(p, [sibling]))

    def test_escape_via_dotdot(self) -> None:
        outside = os.path.join(os.path.dirname(self.root), "escape.txt")
        escaped = os.path.join(self.root, "..", "escape.txt")
        self.assertFalse(local_fs.is_within(escaped, [self.root]))
        self.assertFalse(os.path.exists(outside))


if __name__ == "__main__":
    unittest.main()
